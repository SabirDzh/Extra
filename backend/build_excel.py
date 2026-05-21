import json
from openpyxl import Workbook
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

def create_excel():
    # 1. Load the parsed errors from JSON
    with open("parsed_errors.json", "r", encoding="utf-8") as f:
        errors = json.load(f)

    # 2. Create the workbook and select active sheet
    wb = Workbook()
    ws = wb.active
    ws.title = "Errors"

    # Enable grid lines visibility explicitly
    ws.views.sheetView[0].showGridLines = True

    # 3. Define headers and columns
    headers = ["title", "description", "image", "is_published", "order_index"]
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.value = header
        # Apply header styling for premium look
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid") # Dark blue
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Set columns width
    ws.column_dimensions["A"].width = 50  # title
    ws.column_dimensions["B"].width = 100 # description
    ws.column_dimensions["C"].width = 30  # image
    ws.column_dimensions["D"].width = 15  # is_published
    ws.column_dimensions["E"].width = 15  # order_index

    # Define common styling for cells
    thin_border = Border(
        left=Side(style='thin', color='D3D3D3'),
        right=Side(style='thin', color='D3D3D3'),
        top=Side(style='thin', color='D3D3D3'),
        bottom=Side(style='thin', color='D3D3D3')
    )

    # 4. Fill data and insert images
    for index, item in enumerate(errors, 2):
        # Set row height to accommodate image beautifully
        ws.row_dimensions[index].height = 120

        # Title
        cell_title = ws.cell(row=index, column=1, value=item["title"])
        cell_title.font = Font(name="Calibri", size=11, bold=True)
        cell_title.alignment = Alignment(vertical="center", wrap_text=True)
        cell_title.border = thin_border

        # Description
        cell_desc = ws.cell(row=index, column=2, value=item["description"])
        cell_desc.font = Font(name="Calibri", size=10)
        cell_desc.alignment = Alignment(vertical="center", wrap_text=True)
        cell_desc.border = thin_border

        # Image placeholder text (optional, but empty is cleaner. Image will overlap it anyway)
        cell_img = ws.cell(row=index, column=3, value="")
        cell_img.border = thin_border
        
        # Load and insert image
        img_path = f"Ошибки/{item['image_path']}"
        try:
            img = OpenpyxlImage(img_path)
            # Standardize image sizes if needed, or keep original size.
            # openpyxl handles original size, but we can set dimensions to fit row height (120 pt ~ 160 px)
            # Let's keep original ratio but scale to fit in the cell gracefully
            original_width = img.width
            original_height = img.height
            
            # Row height 120 pt is approx 160 pixels. Let's scale image to height 150px
            target_height = 150
            scale_ratio = target_height / original_height
            img.width = int(original_width * scale_ratio)
            img.height = target_height
            
            ws.add_image(img, f"C{index}")
        except Exception as e:
            print(f"Error inserting image {img_path}: {e}")

        # Is Published
        cell_pub = ws.cell(row=index, column=4, value="TRUE" if item["is_published"] else "FALSE")
        cell_pub.font = Font(name="Calibri", size=11)
        cell_pub.alignment = Alignment(horizontal="center", vertical="center")
        cell_pub.border = thin_border

        # Order Index
        cell_order = ws.cell(row=index, column=5, value=item["order_index"])
        cell_order.font = Font(name="Calibri", size=11)
        cell_order.alignment = Alignment(horizontal="center", vertical="center")
        cell_order.border = thin_border

    # 5. Save the workbook
    output_filename = "errors_import.xlsx"
    wb.save(output_filename)
    print(f"Successfully generated {output_filename} with {len(errors)} records.")

if __name__ == "__main__":
    create_excel()
