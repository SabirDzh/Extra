import csv
import io
import json
import uuid
from typing import List, Literal

import pandas as pd
from core.models.error import Error
from core.schemas.error import ErrorCreate, ErrorUpdate
from fastapi import HTTPException, UploadFile, status
from sqlalchemy import delete, insert, select, func
from sqlalchemy.ext.asyncio import AsyncSession
from utils.db import ensure_unique_field


async def get_errors(
    session: AsyncSession,
    limit: int = 20,
    offset: int = 0,
    sorted: Literal["asc", "desc"] = "asc",
) -> List[Error]:
    stmt = select(Error).offset(offset).limit(limit).order_by(
        Error.order_index.asc() if sorted == "asc" else Error.order_index.desc()
    )
    result = await session.execute(stmt)
    return result.scalars().all()


async def get_error(
    session: AsyncSession,
    error_id: uuid.UUID,
) -> Error | None:
    return await session.get(Error, error_id)


async def search_errors(
    session: AsyncSession,
    q: str | None = None,
    limit: int = 20,
    offset: int = 0,
) -> List[Error]:
    query = select(Error)
    
    if q:
        if len(q) < 2:
            search_pattern = f"%{q}%"
            query = query.where(
                (Error.title.ilike(search_pattern))
                | (Error.description.ilike(search_pattern))
            )
            query = query.order_by(Error.order_index.asc())
        else:
            relevance = (
                func.similarity(Error.title, q) * 2
                + func.similarity(Error.description, q) * 0.5
            )
            search_pattern = f"%{q}%"
            query = query.where(
                (Error.title.bool_op("%")(q))
                | (Error.description.bool_op("%")(q))
                | (Error.title.ilike(search_pattern))
                | (Error.description.ilike(search_pattern))
            )
            query = query.order_by(relevance.desc())
    else:
        query = query.order_by(Error.order_index.asc())

    query = query.offset(offset).limit(limit)
    result = await session.execute(query)
    return result.scalars().all()


async def create_error(
    session: AsyncSession,
    error_in: ErrorCreate,
    user_id: uuid.UUID,
) -> Error:
    await ensure_unique_field(
        session,
        Error,
        "title",
        error_in.title,
        error_msg=f"Error with title '{error_in.title}' already exists",
    )

    error = Error(**error_in.model_dump(), created_by=user_id)
    session.add(error)
    await session.commit()
    await session.refresh(error)
    return error


async def update_error(
    session: AsyncSession,
    error: Error,
    error_update: ErrorUpdate,
) -> Error:
    patch = error_update.model_dump(exclude_unset=True)

    if "title" in patch and patch["title"]:
        await ensure_unique_field(
            session,
            Error,
            "title",
            patch["title"],
            exclude_id=error.id,
            error_msg=f"Error with title '{patch['title']}' already exists",
        )

    for field, value in patch.items():
        setattr(error, field, value)

    await session.commit()
    await session.refresh(error)
    return error


async def delete_error(
    session: AsyncSession,
    error: Error,
) -> None:
    await session.delete(error)
    await session.commit()


async def delete_errors(
    session: AsyncSession,
    error_ids: List[uuid.UUID],
) -> None:
    stmt = delete(Error).where(Error.id.in_(error_ids))
    await session.execute(stmt)
    await session.commit()


async def delete_all_errors(session: AsyncSession) -> None:
    stmt = delete(Error)
    await session.execute(stmt)
    await session.commit()


def check_format(filename: str) -> str:
    filename = filename.lower()
    if not filename.endswith((".csv", ".xlsx", ".xls", ".json")):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Разрешены только файлы форматов JSON, CSV и Excel (.xlsx, .xls)",
        )
    return filename


def parse_error_json_file(contents: bytes) -> list[dict]:
    try:
        items = json.loads(contents.decode("utf-8"))
        if not isinstance(items, list):
            raise ValueError("Root element must be a list")
        return items
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Invalid JSON content: {str(e)}")


def parse_error_csv_file(contents: bytes) -> list[dict]:
    try:
        text = contents.decode("utf-8")
    except UnicodeDecodeError:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Файл CSV должен быть в кодировке UTF-8",
        )

    errors_data = []
    reader = csv.DictReader(io.StringIO(text))
    for row in reader:
        title = row.get("title", "").strip()
        description = row.get("description", "").strip()

        if title and description:
            order_index = 0
            if "order_index" in row and row["order_index"].isdigit():
                order_index = int(row["order_index"])

            is_published = True
            if "is_published" in row:
                val = row["is_published"].strip().lower()
                if val in ["false", "0", "no"]:
                    is_published = False

            errors_data.append(
                {
                    "title": title,
                    "description": description,
                    "order_index": order_index,
                    "is_published": is_published,
                }
            )
    return errors_data


def parse_error_excel_file(contents: bytes) -> list[dict]:
    df = pd.read_excel(io.BytesIO(contents))

    # Support for both English and Russian column names
    title_col = None
    if "title" in df.columns:
        title_col = "title"
    elif "название" in df.columns:
        title_col = "название"

    desc_col = None
    if "description" in df.columns:
        desc_col = "description"
    elif "описание" in df.columns:
        desc_col = "описание"

    if not title_col or not desc_col:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="В файле Excel отсутствуют обязательные колонки 'title'/'название' и/или 'description'/'описание'",
        )

    df = df.fillna("")
    
    try:
        from openpyxl import load_workbook
        from openpyxl_image_loader import SheetImageLoader
        wb = load_workbook(io.BytesIO(contents), data_only=True)
        sheet = wb.active
        image_loader = SheetImageLoader(sheet)
        has_images = True
    except Exception:
        has_images = False
        image_loader = None

    errors_data = []

    for index, row in df.iterrows():
        title = str(row.get(title_col, "")).strip()
        description = str(row.get(desc_col, "")).strip()

        if title and description:
            is_published = True
            published_col = None
            for col in ["is_published", "опубликовано", "опубликован"]:
                if col in df.columns:
                    published_col = col
                    break

            if published_col:
                val = str(row.get(published_col, "")).strip().lower()
                if val in ["false", "0", "no", "ложь", "нет"]:
                    is_published = False

            order_index = 0
            order_col = None
            for col in ["order_index", "порядок"]:
                if col in df.columns:
                    order_col = col
                    break

            if order_col:
                val = row.get(order_col, "")
                if str(val).isdigit() or isinstance(val, (int, float)):
                    try:
                        order_index = int(val)
                    except ValueError:
                        pass

            img_obj = None
            if has_images:
                excel_row = index + 2
                for col_idx in range(1, len(df.columns) + 1):
                    from openpyxl.utils import get_column_letter
                    col_letter = get_column_letter(col_idx)
                    cell_coord = f"{col_letter}{excel_row}"
                    if image_loader and image_loader.image_in(cell_coord):
                        try:
                            img_obj = image_loader.get(cell_coord)
                            break
                        except Exception:
                            pass

            errors_data.append(
                {
                    "title": title,
                    "description": description,
                    "order_index": order_index,
                    "is_published": is_published,
                    "image": img_obj,
                }
            )

    return errors_data


async def import_errors(
    session: AsyncSession,
    file: UploadFile,
    user_id: uuid.UUID,
) -> dict:
    filename = check_format(file.filename)

    try:
        contents = await file.read()

        if filename.endswith(".json"):
            items = parse_error_json_file(contents)
        elif filename.endswith(".csv"):
            items = parse_error_csv_file(contents)
        else:
            items = parse_error_excel_file(contents)

        if not items:
            return {"message": "Файл пуст или не содержит валидных данных для импорта"}

        # Exclude duplicates
        incoming_titles = []
        valid_items = []
        for item in items:
            if "title" in item and "description" in item:
                t = item["title"].strip()
                if t:
                    incoming_titles.append(t)
                    valid_items.append(item)

        if not valid_items:
            return {"message": "Не найдено валидных заголовков в файле."}

        stmt = select(Error.title).where(Error.title.in_(incoming_titles))
        existing_result = await session.execute(stmt)
        existing_titles = {t.lower() for t in existing_result.scalars().all()}

        import os
        import uuid
        media_dir = "media/error_img"
        os.makedirs(media_dir, exist_ok=True)
        
        new_errors = []
        
        for item in valid_items:
            if item["title"].lower() not in existing_titles:
                new_error_data = {
                    "title": item["title"],
                    "description": item["description"],
                    "order_index": item.get("order_index", 0),
                    "is_published": item.get("is_published", True),
                    "created_by": user_id,
                    "image": None,
                }
                existing_titles.add(item["title"].lower())
                
                img_obj = item.get("image")
                if img_obj:
                    try:
                        filename_webp = f"{uuid.uuid4()}.webp"
                        filepath = os.path.join(media_dir, filename_webp)
                        if img_obj.mode not in ('RGB', 'RGBA'):
                            img_obj = img_obj.convert('RGBA')
                        img_obj.save(filepath, "WEBP")
                        
                        new_error_data["image"] = f"/{media_dir}/{filename_webp}"
                    except Exception as e:
                        print(f"Error saving image for {item['title']}: {e}")
                
                new_errors.append(new_error_data)

        if not new_errors:
            return {
                "message": f"Найдено {len(valid_items)} записей, но все они уже существуют."
            }

        stmt_insert = insert(Error).values(new_errors)
        await session.execute(stmt_insert)
        await session.commit()

        return {
            "msg": f"Успешно импортировано ошибок: {len(new_errors)}. Пропущено дубликатов: {len(valid_items) - len(new_errors)}.",
            "status": "OK",
        }

    except HTTPException:
        raise
    except Exception as e:
        await session.rollback()
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Ошибка при импорте: {str(e)}",
        )
